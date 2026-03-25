from docx import Document
import re

def inspect_docx(file_path):
    try:
        doc = Document(file_path)
        print(f"Документ содержит {len(doc.paragraphs)} абзацев")
        
        # Проверяем наличие гиперссылок в документе
        hyperlinks_dict = {}
        for rel_id, rel in doc.part.rels.items():
            if "hyperlink" in rel.reltype:
                hyperlinks_dict[rel_id] = rel._target
                print(f"Обнаружена гиперссылка [ID: {rel_id}]: {rel._target}")
        
        print(f"Документ содержит {len(hyperlinks_dict)} гиперссылок.")
        
        # Анализ содержимого абзацев
        for i, para in enumerate(doc.paragraphs):
            text = para.text.strip()
            if text:
                # Проверка на различные форматы ссылок
                is_telegram_link = 'https://t.me/' in text or 't.me/' in text
                has_url = bool(re.search(r'https?://\S+', text))
                has_zakon_reference = 'zakon' in text.lower() or 'закон' in text.lower()
                
                # Попытка определить структуру: заголовок + ссылка или только заголовок
                parts = text.split()
                last_part = parts[-1] if parts else ""
                is_likely_url = last_part.startswith(('http', 'www', 't.me')) if parts else False
                
                link_type = "НЕИЗВЕСТНО"
                if is_telegram_link:
                    link_type = "TELEGRAM"
                elif has_url:
                    link_type = "URL"
                elif has_zakon_reference:
                    link_type = "ZAKON_REFERENCE"
                else:
                    link_type = "ТЕКСТ"
                
                print(f"Абзац {i+1}: [{link_type}] {text[:100]}{'...' if len(text) > 100 else ''}")
                
                # Проверяем наличие прямых связей с гиперссылками
                found_hyperlinks = []
                for run in para.runs:
                    for key, val in run._element.items():
                        if '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}id' in key:
                            rel_id = val
                            if rel_id in hyperlinks_dict:
                                found_hyperlinks.append((rel_id, hyperlinks_dict[rel_id]))
                                
                if found_hyperlinks:
                    print(f"  - НАЙДЕНЫ ПРЯМЫЕ СВЯЗИ С ГИПЕРССЫЛКАМИ:")
                    for rel_id, url in found_hyperlinks:
                        print(f"  - Связь [ID: {rel_id}]: {url}")
                
                # Если это потенциально комбинация заголовка и ссылки, анализируем подробнее
                if has_url:
                    print(f"  - Возможная структура: заголовок + ссылка")
                    # Находим все URL в тексте
                    urls = re.findall(r'https?://\S+', text)
                    for url in urls:
                        print(f"  - Найден URL: {url}")
                    
                    # Пытаемся разделить заголовок и ссылку
                    url_start = text.find('http')
                    if url_start > 0:
                        title = text[:url_start].strip()
                        url = text[url_start:].strip()
                        print(f"  - Заголовок: {title}")
                        print(f"  - Ссылка: {url}")
                
                # Проверяем наличие номеров сообщений Telegram
                tg_msg_ids = re.findall(r't\.me/[^/]+/(\d+)', text)
                if tg_msg_ids:
                    print(f"  - Найдены ID сообщений Telegram: {', '.join(tg_msg_ids)}")
    
    except Exception as e:
        print(f"Ошибка при анализе документа: {str(e)}")

def analyze_hyperlinks(file_path):
    """Специальный анализ гиперссылок и их привязки к абзацам"""
    try:
        doc = Document(file_path)
        print("\n=== ДЕТАЛЬНЫЙ АНАЛИЗ ГИПЕРССЫЛОК ===\n")
        
        # Получаем все гиперссылки
        hyperlinks_dict = {}
        for rel_id, rel in doc.part.rels.items():
            if "hyperlink" in rel.reltype:
                hyperlinks_dict[rel_id] = rel._target
        
        # Словарь для отслеживания использования ссылок
        hyperlinks_usage = {rel_id: [] for rel_id in hyperlinks_dict}
        
        # Анализируем каждый абзац
        for i, para in enumerate(doc.paragraphs):
            text = para.text.strip()
            if not text:
                continue
                
            found_links = False
            
            # Проверяем каждый run на наличие гиперссылок
            for j, run in enumerate(para.runs):
                for key, val in run._element.items():
                    if '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}id' in key:
                        rel_id = val
                        if rel_id in hyperlinks_dict:
                            hyperlinks_usage[rel_id].append((i+1, text))
                            found_links = True
                            print(f"Абзац {i+1}, фрагмент {j+1}: связан с гиперссылкой [ID: {rel_id}]")
                            print(f"  - Текст абзаца: {text[:100]}{'...' if len(text) > 100 else ''}")
                            print(f"  - Текст фрагмента: {run.text[:50]}{'...' if len(run.text) > 50 else ''}")
                            print(f"  - Ссылка: {hyperlinks_dict[rel_id]}")
            
            if not found_links:
                print(f"Абзац {i+1}: не содержит гиперссылок")
                print(f"  - Текст: {text[:100]}{'...' if len(text) > 100 else ''}")
        
        print("\n=== СТАТИСТИКА ИСПОЛЬЗОВАНИЯ ГИПЕРССЫЛОК ===\n")
        for rel_id, usages in hyperlinks_usage.items():
            print(f"Гиперссылка [ID: {rel_id}]: {hyperlinks_dict[rel_id]}")
            if usages:
                print(f"  - Использована в {len(usages)} абзацах: {', '.join([str(p[0]) for p in usages])}")
            else:
                print("  - Не использована в тексте")
        
    except Exception as e:
        print(f"Ошибка при анализе гиперссылок: {str(e)}")

if __name__ == "__main__":
    file_path = "news_for_test.docx"
    inspect_docx(file_path)
    analyze_hyperlinks(file_path) 